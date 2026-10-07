"""10+ slaydli taqdimot kengaytmasi va bo‘sh rasm so‘rovi regressiyasi."""
from io import BytesIO
from pathlib import Path
import sys

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import utils


def test_builders_expand_content_without_detached_slide_error():
    """T33/T34 10 kontent slaydida xulosa oxirida qoladi va saqlanadi."""
    for template_num, builder in (
        (33, utils.build_slide_structure_33),
        (34, utils.build_slide_structure_34),
    ):
        presentation = Presentation(ROOT / "templates" / "shablonlar" / f"{template_num}.pptx")
        builder(presentation, requested_slide_count=10)

        # Muqova + reja + 10 kontent + xulosa.
        assert len(presentation.slides) == 13
        assert len(list(presentation.slides)[2:-1]) == 10
        output = BytesIO()
        presentation.save(output)
        assert len(output.getvalue()) > 1_000


def test_template_33_generates_ten_slides_when_image_query_is_null():
    """Kontent API image_query=null yuborsa ham 10-slaydli fayl yaratiladi."""
    original_fetch_image = utils.fetch_image
    utils.fetch_image = lambda *_args, **_kwargs: None
    try:
        presentation = Presentation(ROOT / "templates" / "shablonlar" / "33.pptx")
        content = [
            {
                "title": f"Bo‘lim {index + 1}",
                "content": ["Sinov matni."],
                "image_query": None,
            }
            for index in range(10)
        ]
        output = utils.generate_template_33_presentation(
            prs=presentation,
            topic="Sinov mavzusi",
            requested_slide_count=10,
            language="uz",
            name_surname="",
            plan={"title": "Reja", "content": ["Kirish", "Xulosa"]},
            content_data_list=content,
            user_images=[],
        )
        rendered = Presentation(BytesIO(output))
        assert len(rendered.slides) == 13
    finally:
        utils.fetch_image = original_fetch_image


def test_empty_image_query_is_not_sent_to_providers():
    """None/bo‘sh rasm so‘rovi tashqi providerga chiqmasdan xavfsiz tugaydi."""
    original_stock = utils._fetch_gold_pixabay_image
    original_deepinfra = utils.fetch_image_deepinfra
    try:
        utils._fetch_gold_pixabay_image = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("stock provider chaqirilmasligi kerak")
        )
        utils.fetch_image_deepinfra = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("AI provider chaqirilmasligi kerak")
        )
        assert utils.fetch_image(None) is None
        assert utils.fetch_image_preview_urls(None) == []
    finally:
        utils._fetch_gold_pixabay_image = original_stock
        utils.fetch_image_deepinfra = original_deepinfra


if __name__ == "__main__":
    test_builders_expand_content_without_detached_slide_error()
    test_template_33_generates_ten_slides_when_image_query_is_null()
    test_empty_image_query_is_not_sent_to_providers()
    print("PRESENTATION_EXPANSION_GUARD_TEST_OK")
